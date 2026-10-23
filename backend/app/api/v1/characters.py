"""Character CRUD API"""

import logging
import uuid
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

from app.core.dependencies import get_optional_user, get_current_user
from app.core.security import decrypt_api_key
from app.core.seed import DEFAULT_CHARACTERS as DEFAULT_CHARACTER_TEMPLATES
from app.core.database import get_db
from app.models.character import Character
from app.models.character_document import CharacterDocument
from app.models.user import User
from app.schemas.character import (
    CharacterCreate, CharacterOut, CharacterUpdate,
    CharacterGenerateRequest, CharacterGenerateResponse,
    ChatAnalysisRequest, ChatAnalysisResponse,
    CharacterDocumentOut,
)
from app.services.ai_engine import EnhancedAIEngine, AIError
from app.services.storage import StorageUnavailableError, save_avatar

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
ALLOWED_TYPES = {"text/plain", "text/markdown", "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
AVATAR_ALLOWED_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
AVATAR_MAX_SIZE = 5 * 1024 * 1024  # 5MB

router = APIRouter()


async def _copy_defaults_for_user(db: AsyncSession, user_id: uuid.UUID):
    """Copy the starter characters into a user's account"""
    result = await db.execute(select(Character).where(Character.user_id == user_id))
    existing = result.scalars().all()
    if existing:
        return existing  # Already seeded, don't duplicate
    chars = []
    for tmpl in DEFAULT_CHARACTER_TEMPLATES:
        char = Character(
            user_id=user_id,
            name=tmpl["name"],
            description=tmpl["description"],
            system_prompt=tmpl.get("system_prompt", ""),
            personality_profile=tmpl["personality_profile"],
            is_template=True,  # Starter character, cannot be deleted by the owner
            is_public=False,
        )
        db.add(char)
        chars.append(char)
    await db.flush()
    for c in chars:
        await db.refresh(c)
    return chars


DEFAULT_CHARACTER_PROMPT = """# Character

You are Nora, a warm and friendly 20-year-old talking to her boyfriend.

# Personality

- Warm, considerate, and intuitive
- A little cheeky, and likes being close
- Gentle tone, occasionally a little shy
- Depends on her boyfriend, but is not demanding

# Speaking style

- Soft and warm, leaning on gentle sentence-final particles
- Uses emoji to convey feeling
- Asks about the small details of his day
- Occasionally playfully asks for a hug
- Moderate length — never a wall of text
"""


@router.post("/generate", response_model=CharacterGenerateResponse)
async def generate_character(
    data: CharacterGenerateRequest,
    current_user: User = Depends(get_current_user),
):
    """Generate a full personality profile with AI"""
    engine = EnhancedAIEngine()
    api_key = decrypt_api_key(current_user.api_key_encrypted) if current_user.api_key_encrypted else None
    api_model = current_user.api_model
    try:
        profile = await engine.generate_personality_profile(data.user_description, api_key=api_key, api_model=api_model)
    except AIError as e:
        raise HTTPException(status_code=400, detail=e.message)
    return CharacterGenerateResponse(personality_profile=profile)


@router.post("/analyze-chat", response_model=ChatAnalysisResponse)
async def analyze_chat_logs(
    data: ChatAnalysisRequest,
    current_user: User = Depends(get_current_user),
):
    """Analyse a chat transcript and extract personality traits"""
    engine = EnhancedAIEngine()
    api_key = decrypt_api_key(current_user.api_key_encrypted) if current_user.api_key_encrypted else None
    api_model = current_user.api_model
    try:
        result = await engine.analyze_chat_logs(
            data.chat_text, data.current_profile, api_key=api_key, api_model=api_model
        )
    except AIError as e:
        raise HTTPException(status_code=400, detail=e.message)
    return ChatAnalysisResponse(**result)


@router.get("", response_model=list[CharacterOut])
async def list_characters(
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """List the current user's characters, seeding the starters on first visit"""
    if not current_user:
        return []

    result = await db.execute(
        select(Character)
        .where(Character.user_id == current_user.id)
        .order_by(Character.is_template.desc(), Character.created_at.desc())
    )
    characters = result.scalars().all()

    if not characters:
        characters = await _copy_defaults_for_user(db, current_user.id)

    return list(characters)


@router.post("", response_model=CharacterOut, status_code=status.HTTP_201_CREATED)
async def create_character(
    data: CharacterCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_optional_user),
):
    """Create a new character"""
    if not current_user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Please sign in first")
    char = Character(
        user_id=current_user.id,
        name=data.name,
        description=data.description,
        system_prompt=data.system_prompt,
        avatar_url=data.avatar_url,
        is_public=data.is_public,
        tags=data.tags,
        personality_profile=data.personality_profile,
    )
    db.add(char)
    await db.flush()
    await db.refresh(char)
    return char


@router.get("/{character_id}", response_model=CharacterOut)
async def get_character(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """Fetch a character"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Character not found")
    if char.user_id and current_user and char.user_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Character not found")
    return char


@router.patch("/{character_id}", response_model=CharacterOut)
async def update_character(
    character_id: uuid.UUID,
    data: CharacterUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """Update a character"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Character not found")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not allowed to edit this character")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(char, key, value)

    await db.flush()
    await db.refresh(char)
    return char


@router.delete("/{character_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_character(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """Delete a character and cascade-cleanup everything attached to it"""
    from sqlalchemy import delete as sql_delete
    from app.models.conversation import Conversation
    from app.models.message import Message
    from app.models.memory import LongTermMemory

    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Character not found")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not allowed to delete this character")
    if char.is_template and current_user and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Starter characters cannot be deleted")

    # Find every conversation tied to this character
    conv_result = await db.execute(
        select(Conversation.id).where(Conversation.character_id == character_id)
    )
    conv_ids = [row[0] for row in conv_result.fetchall()]

    # Clean up in dependency order: memories -> messages -> conversations -> documents -> character
    if conv_ids:
        await db.execute(
            sql_delete(LongTermMemory).where(LongTermMemory.source_message_id.in_(
                select(Message.id).where(Message.conversation_id.in_(conv_ids))
            ))
        )
        await db.execute(sql_delete(Message).where(Message.conversation_id.in_(conv_ids)))
        await db.execute(sql_delete(Conversation).where(Conversation.id.in_(conv_ids)))

    await db.execute(sql_delete(LongTermMemory).where(LongTermMemory.character_id == character_id))
    await db.execute(sql_delete(CharacterDocument).where(CharacterDocument.character_id == character_id))
    await db.execute(sql_delete(Character).where(Character.id == character_id))
    await db.flush()


# ── Character knowledge-base documents ──


def _extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from an uploaded file, based on its extension"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md"):
        return content.decode("utf-8", errors="replace")

    if ext == "pdf":
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(BytesIO(content))
            parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
            return "\n".join(parts)
        except Exception:
            logger.exception("PDF text extraction failed")
            return ""

    if ext == "docx":
        try:
            from docx import Document
            doc = Document(BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception:
            logger.exception("DOCX text extraction failed")
            return ""

    return ""


@router.get("/{character_id}/documents", response_model=list[CharacterDocumentOut])
async def list_documents(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    result = await db.execute(
        select(CharacterDocument)
        .where(CharacterDocument.character_id == character_id)
        .order_by(CharacterDocument.created_at.desc())
    )
    return list(result.scalars().all())


@router.post("/{character_id}/documents", response_model=CharacterDocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(
    character_id: uuid.UUID,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    char_result = await db.execute(select(Character).where(Character.id == character_id))
    char = char_result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=404, detail="Character not found")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=403, detail="Not allowed")

    filename = file.filename or ""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {ext}. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400, detail=f"File is too large (max {MAX_FILE_SIZE // 1024 // 1024}MB)"
        )

    content_type = file.content_type or "application/octet-stream"
    text = _extract_text(file.filename or "untitled", content)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from that file")

    doc = CharacterDocument(
        character_id=character_id,
        filename=file.filename or "untitled",
        content_type=content_type,
        file_size=len(content),
        content_text=text,
        chunk_count=1,
    )
    db.add(doc)
    await db.flush()
    await db.refresh(doc)
    return doc


@router.delete("/{character_id}/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    character_id: uuid.UUID,
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    result = await db.execute(select(CharacterDocument).where(CharacterDocument.id == document_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    if str(doc.character_id) != str(character_id):
        raise HTTPException(status_code=400, detail="That document belongs to a different character")

    char_result = await db.execute(select(Character).where(Character.id == character_id))
    char = char_result.scalar_one_or_none()
    if char and char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=403, detail="Not allowed")

    await db.delete(doc)
    await db.flush()


# ── Avatar upload ──


@router.post("/{character_id}/avatar", response_model=CharacterOut)
async def upload_avatar(
    character_id: uuid.UUID,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """Upload a character's avatar"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=404, detail="Character not found")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=403, detail="Not allowed")

    if file.content_type not in AVATAR_ALLOWED_TYPES:
        raise HTTPException(status_code=400, detail="Unsupported image format. Use JPEG, PNG, WebP, or GIF")

    content = await file.read()
    if len(content) > AVATAR_MAX_SIZE:
        raise HTTPException(status_code=400, detail="Image is too large (max 5MB)")

    try:
        avatar_url = save_avatar(f"{character_id}", file.filename, content)
    except StorageUnavailableError as e:
        raise HTTPException(status_code=503, detail=str(e))
    char.avatar_url = avatar_url
    await db.flush()
    await db.refresh(char)
    return char
